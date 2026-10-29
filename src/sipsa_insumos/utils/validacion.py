"""Validación de cierre de un período — SIPSA Insumos.

Sin SAS como referencia, cada mes se valida la consistencia del propio proceso.
`validar_periodo` revisa:

  1. Configuración  — globals y rutas de los módulos activos apuntan al período.
  2. Insumos        — existen la base liviana y las DIVIPOLA; versión usada
                      (inicial / ajustada) y checksum frente a OneDrive.
  3. Ejecución      — las salidas son posteriores a los insumos que usan y
                      salieron de la base vigente (mismo número de filas); si
                      hubo ejecución definitiva, el respaldo preliminar existe
                      y la base ajustada sigue en raw.
  4. Completitud    — cada módulo genera los mismos archivos que en su período
                      anterior; serie departamental (5 archivos) y Sin Precio
                      Anterior (módulos activos).
  5. Diagnósticos   — los FALTAN vacíos; filas en DUPLICADOS como alerta.
  6. Continuidad    — los archivos del mes anterior que llegan con el período
                      (MAYORESQUE2 y "Revisión …") son idénticos a la salida
                      Python de ese mes.

Cada verificación queda en "ok", "alerta" o "error". El resultado se guarda en
data/08_reporting/<P>/VALIDACION_<P>.json y VALIDACION_<P>.xlsx.

Uso: python scripts/validar_periodo.py [PERIODO]   (por defecto, el de globals.yml)
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import math
import re
import unicodedata
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from python_calamine import CalamineWorkbook

from sipsa_insumos.utils.insumos import ruta_vigente, version_insumo

_MESES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]
_NOMBRES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
            "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

# Calendario de módulos (igual al de app.py) y carpeta de la serie departamental
_CALENDARIO = {
    "agricolas": range(1, 13), "pecuarios": range(1, 13),
    "elementos": [1, 3, 5, 7, 9, 11], "empaques": [1, 3, 5, 7, 9, 11],
    "propagacion": [2, 4, 6, 8, 10, 12],
    "arriendos": [2, 5, 8, 11], "servicios": [2, 5, 8, 11],
    "jornales": [3, 6, 9, 12], "especies": [3, 6, 9, 12],
}
_SALTO = {"agricolas": 1, "pecuarios": 1, "elementos": 2, "empaques": 2, "propagacion": 2,
          "arriendos": 3, "servicios": 3, "jornales": 3, "especies": 3}
_CARPETA_DEPTAL = {"agricolas": "Ins_Agrícolas", "pecuarios": "Ins_Pecuarios", "elementos": "Elementos",
                   "empaques": "Empaques", "propagacion": "Material_Propaga", "arriendos": "Arriendos",
                   "servicios": "Servicios", "jornales": "Jornales", "especies": "Especies"}
_CLAVE = {"agricolas": "agrícolas", "pecuarios": "pecuarios", "elementos": "elementos", "empaques": "empaques",
          "propagacion": "propagación", "arriendos": "arriendos", "servicios": "servicios",
          "jornales": "jornales", "especies": "especie"}
_CALENDARIO_SPA = {"agricolas": range(1, 13), "pecuarios": range(1, 13),
                   "elementos": [1, 3, 5, 7, 9, 11], "propagacion": [2, 4, 6, 8, 10, 12]}
_NUMERO = re.compile(r"^-?\d+(\.\d+)?$")
_TOKEN_PERIODO = re.compile(r"(ENE|FEB|MAR|ABR|MAY|JUN|JUL|AGO|AUG|SEP|OCT|NOV|DIC)\d{4}", re.IGNORECASE)


# ── utilidades ───────────────────────────────────────────────────────────────

def _nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s)


def _mes(periodo: str) -> int:
    return _MESES.index(periodo[:3].upper()) + 1


def _anterior(periodo: str, salto: int) -> str:
    m, a = _mes(periodo) - salto, int(periodo[3:])
    while m < 1:
        m, a = m + 12, a - 1
    return f"{_MESES[m - 1]}{a}"


def _excel(d: Path) -> list[Path]:
    return sorted(f for f in d.glob("*.xls*") if f.is_file() and not f.name.startswith("~$")) if d.is_dir() else []


def _plantilla(nombre: str) -> str:
    """Nombre de archivo sin el período (para comparar meses): 'X_SEP2026.xlsx' → 'X_<P>'."""
    return _TOKEN_PERIODO.sub("<P>", _nfc(Path(nombre).stem)).upper()


def _llegada(f: Path) -> float:
    """Momento en que el insumo quedó en raw. Si la fecha de modificación está en el
    futuro (archivo de un equipo con el reloj adelantado), se usa la de creación."""
    st = f.stat()
    creado = getattr(st, "st_birthtime", st.st_ctime)
    return creado if st.st_mtime > dt.datetime.now().timestamp() else max(creado, st.st_mtime)


def _md5(f: Path) -> str:
    return hashlib.md5(f.read_bytes()).hexdigest()


def _hojas(f: Path) -> dict[str, list[list]]:
    wb = CalamineWorkbook.from_path(str(f))
    return {s: wb.get_sheet_by_name(s).to_python() for s in wb.sheet_names}


def _filas_con_datos(f: Path) -> int:
    hojas = _hojas(f)
    filas = next(iter(hojas.values()), [])
    return sum(1 for r in filas[1:] if any(v not in ("", None) for v in r))


def _filas_cuerpo(f: Path) -> int:
    return sum(max(len(h) - 1, 0) for h in _hojas(f).values())


def _celda(v):
    """Valor comparable: un número guardado como texto ('34662') vale lo mismo que
    el número; los decimales se comparan con 12 cifras significativas."""
    if isinstance(v, str):
        v = v.strip()
        if _NUMERO.match(v):
            v = float(v)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(f"{float(v):.12g}") if math.isfinite(v) else None
    return None if v in ("", None) else v


def _clave_fila(t: tuple) -> tuple:
    return tuple((x is None, type(x).__name__, x if x is not None else 0) for x in t)


def _mismo_contenido(f1: Path, f2: Path) -> tuple[str, str]:
    """Compara dos Excel por valores: 'igual', 'orden' (mismas filas, otro orden) o 'distinto'."""
    h1, h2 = _hojas(f1), _hojas(f2)
    if len(h1) != len(h2):
        return "distinto", f"{len(h1)} vs {len(h2)} hojas"
    estado = "igual"
    for (n1, r1), (_, r2) in zip(h1.items(), h2.items()):
        a = [tuple(_celda(v) for v in fila) for fila in r1]
        b = [tuple(_celda(v) for v in fila) for fila in r2]
        if a[:1] != b[:1]:
            return "distinto", f"hoja '{n1}': columnas distintas"
        if len(a) != len(b):
            return "distinto", f"hoja '{n1}': {len(a) - 1} vs {len(b) - 1} filas"
        if a != b:
            if sorted(a[1:], key=_clave_fila) != sorted(b[1:], key=_clave_fila):
                i = next(i for i, (x, y) in enumerate(zip(a, b)) if x != y)
                return "distinto", f"hoja '{n1}': valores distintos desde la fila {i + 1}"
            estado = "orden"
    return estado, "idéntico" if estado == "igual" else "mismas filas, en otro orden"


def _leer_yaml(p: Path) -> dict:
    return yaml.safe_load(p.read_text(encoding="utf-8")) or {} if p.exists() else {}


# ── validación ───────────────────────────────────────────────────────────────

class _Resultado:
    def __init__(self):
        self.items: list[dict] = []

    def add(self, grupo: str, modulo: str, verificacion: str, estado: str, detalle: str = ""):
        self.items.append({"grupo": grupo, "modulo": modulo, "verificacion": verificacion,
                           "estado": estado, "detalle": detalle})


def validar_periodo(periodo: str | None = None, root: str | Path = ".", guardar: bool = True,
                    ruta_onedrive: str | None = None) -> dict:
    root = Path(root)
    conf = root / "conf" / "base"
    glob_ = _leer_yaml(conf / "globals.yml")
    periodo = (periodo or str(glob_.get("periodo", ""))).upper()
    raw, rep = root / "data" / "01_raw" / periodo, root / "data" / "08_reporting" / periodo
    mes = _mes(periodo)
    activos = [m for m, meses in _CALENDARIO.items() if mes in meses]
    es_actual = str(glob_.get("periodo", "")).upper() == periodo
    spa = _leer_yaml(conf / "parameters_sin_precio_ant.yml").get("sin_precio_ant", {})
    if es_actual:   # en el período actual se respeta lo que el usuario activó en SPA
        spa_activos = [m for m, c in spa.items() if isinstance(c, dict) and c.get("activo")]
    else:
        spa_activos = [m for m, meses in _CALENDARIO_SPA.items() if mes in meses]
    conf_val = _leer_yaml(conf / "parameters.yml").get("validacion") or {}
    ruta_onedrive = ruta_onedrive or conf_val.get("ruta_onedrive_insumos")
    excepciones = {_nfc(e["archivo"]).lower(): e.get("motivo", "")
                   for e in conf_val.get("diferencias_conocidas") or []}
    r = _Resultado()

    # 1. Configuración ─────────────────────────────────────────────────────────
    if es_actual:
        r.add("Configuración", "—", "Período configurado", "ok", f"globals.yml = {periodo}")
    else:
        r.add("Configuración", "—", "Período configurado", "ok",
              f"se valida un período anterior ({periodo}); el configurado es {glob_.get('periodo')}")
    for m in activos if es_actual else []:
        p = _leer_yaml(conf / f"parameters_{m}.yml").get(m, {})
        malas = [f"{c}: {p.get(c)}" for c in ("archivo_liviana", "archivo_divipola_grupos")
                 if p.get(c) and f"/{periodo}/" not in str(p.get(c)).replace("\\", "/")]
        r.add("Configuración", m, "Rutas de insumos del período", "error" if malas else "ok",
              "; ".join(malas) if malas else "base liviana y DIVIPOLA apuntan al período")

    # 2. Insumos ───────────────────────────────────────────────────────────────
    carpeta_od = None
    if ruta_onedrive and Path(ruta_onedrive).is_dir():
        carpeta_od = next((d for d in Path(ruta_onedrive).iterdir() if d.is_dir()
                           and _nfc(d.name).upper().startswith(f"{mes}. {_NOMBRES[mes - 1]}")
                           and d.name.endswith(periodo[3:])), None)
    if not carpeta_od:
        r.add("Insumos", "—", "Carpeta del período en OneDrive", "alerta",
              "no se encontró; no se comparan checksums")
    maestro = ruta_vigente(raw / f"DIVIPOLA {periodo}" / "DIVIPOLA.xlsx")
    r.add("Insumos", "—", "DIVIPOLA maestro", "ok" if maestro.exists() else "error",
          f"{version_insumo(maestro)}: {maestro.name}" if maestro.exists() else "no existe")
    bases: dict[str, Path] = {}
    for m in activos:
        for sub, prefijo, etiqueta in ((f"BASES LIVIANAS {periodo}", "", "Base liviana"),
                                       (f"DIVIPOLA {periodo}", "divipola", "DIVIPOLA del módulo")):
            carpeta = raw / sub
            cands = [f for f in _excel(carpeta) if _CLAVE[m] in _nfc(f.name).lower()
                     and _nfc(f.name).lower().startswith(prefijo)]
            if len(cands) != 1:
                r.add("Insumos", m, etiqueta, "error",
                      f"se esperaba 1 archivo con '{_CLAVE[m]}' en {carpeta}; hay {len(cands)}")
                continue
            inicial = cands[0]
            vigente = ruta_vigente(inicial)
            if etiqueta == "Base liviana":
                bases[m] = vigente
            detalle = f"versión {version_insumo(vigente)}: {vigente.name}"
            estado = "ok"
            if carpeta_od:
                od = carpeta_od / inicial.parent.name / inicial.name
                if not od.exists():
                    estado, detalle = "alerta", detalle + " | no está en OneDrive"
                elif _md5(od) == _md5(inicial):
                    detalle += " | inicial igual a OneDrive"
                elif _mismo_contenido(od, inicial)[0] != "distinto":
                    detalle += " | inicial con el mismo contenido que OneDrive (solo cambia el archivo)"
                else:
                    estado, detalle = "error", detalle + " | inicial DISTINTA a la de OneDrive"
            r.add("Insumos", m, etiqueta, estado, detalle)

    # 3. Ejecución y 4. completitud por módulo ─────────────────────────────────
    for m in activos:
        salidas = _excel(rep / m)
        if not salidas:
            r.add("Completitud", m, "Reportes municipales", "error", f"no hay salidas en {rep / m}")
            continue
        anterior = _anterior(periodo, _SALTO[m])
        ref = _excel(root / "data" / "08_reporting" / anterior / m)
        if ref:
            faltan = sorted({_plantilla(f.name) for f in ref} - {_plantilla(f.name) for f in salidas})
            r.add("Completitud", m, "Reportes municipales", "error" if faltan else "ok",
                  f"faltan respecto a {anterior}: {', '.join(faltan)}" if faltan
                  else f"{len(salidas)} archivos, los mismos que en {anterior}")
        else:
            r.add("Completitud", m, "Reportes municipales", "alerta",
                  f"{len(salidas)} archivos; no hay salida de {anterior} para comparar")

        base = bases.get(m)
        if base and base.exists():
            mas_vieja = min(f.stat().st_mtime for f in salidas)
            insumos = [base] + [x for x in [maestro] if x.exists()]
            nuevos = [x.name for x in insumos if _llegada(x) > mas_vieja]
            r.add("Ejecución", m, "Salidas posteriores a los insumos", "error" if nuevos else "ok",
                  f"hay insumos más nuevos que las salidas ({', '.join(nuevos)}): falta volver a correr"
                  if nuevos else f"generadas el {dt.datetime.fromtimestamp(mas_vieja):%d/%m/%Y %H:%M}")
            arch_base = next((f for f in salidas if f.name.upper().startswith("BASE_")), None)
            if arch_base:
                n_sal, n_base = _filas_cuerpo(arch_base), _filas_con_datos(base)
                r.add("Ejecución", m, "Salidas hechas con la base vigente", "ok" if n_sal == n_base else "error",
                      f"{arch_base.name}: {n_sal} filas = base {version_insumo(base)} ({n_base})" if n_sal == n_base
                      else f"{arch_base.name} tiene {n_sal} filas y la base {version_insumo(base)} {n_base}")
            if (rep / "_PRELIMINAR" / m).is_dir() and version_insumo(base) == "inicial":
                r.add("Ejecución", m, "Base ajustada disponible", "error",
                      "hubo ejecución definitiva pero la base ajustada ya no está en "
                      f"AJUSTADOS {periodo}/: una nueva corrida usaría la inicial")
            elif version_insumo(base) == "ajustada" and not (rep / "_PRELIMINAR" / m).is_dir():
                r.add("Ejecución", m, "Respaldo de la versión preliminar", "alerta",
                      f"no existe {rep / '_PRELIMINAR' / m}")

        deptal = _excel(rep / "serie_deptal" / _CARPETA_DEPTAL[m])
        r.add("Completitud", m, "Serie departamental", "ok" if len(deptal) == 5 else "error",
              f"{len(deptal)} de 5 archivos")

    for m in spa_activos:
        nombre = (spa[m].get("formato_sas") or {}).get("nombre_salida", m.upper())
        f = rep / "sin_precio_ant" / f"REV_SIN_PRECIO_ANTE_{nombre}_{periodo}.xlsx"
        r.add("Completitud", m, "Sin Precio Anterior", "ok" if f.exists() else "error",
              f.name if f.exists() else f"no existe {f.name}")
    for m in [x for x in _CALENDARIO if x not in activos and (rep / x).is_dir()]:
        r.add("Completitud", m, "Módulo fuera de calendario", "alerta",
              f"hay salidas de {m} aunque no corresponde en {_MESES[mes - 1]}")

    # 5. Diagnósticos ─────────────────────────────────────────────────────────
    for f in sorted(rep.rglob("*.xls*")):
        if "_PRELIMINAR" in f.parts or f.name.startswith("~$"):
            continue
        nombre = f.name.upper()
        modulo = f.parent.name if f.parent != rep else "—"
        if nombre.startswith("FALTAN"):
            n = _filas_cuerpo(f)
            r.add("Diagnósticos", modulo, f.name, "error" if n else "ok",
                  f"{n} registros sin grupo / nombre de publicación: revisar la DIVIPOLA" if n else "vacío")
        elif "DUPLI" in nombre:
            n = _filas_cuerpo(f)
            if n:
                r.add("Diagnósticos", modulo, f.name, "alerta", f"{n} registros duplicados")

    # 6. Continuidad con el mes anterior ───────────────────────────────────────
    recibidos = [f for f in _excel(raw) if "MAYORESQUE2" in f.name.upper()]
    recibidos += [f for d in raw.glob("REVISI*") for f in _excel(d) if f.name.lower().startswith("revisi")]
    reporting = root / "data" / "08_reporting"
    for f in recibidos:
        candidatos = [g for g in reporting.glob("*/*/*.xls*")
                      if g.parent.parent.name not in (periodo, "_PRELIMINAR")
                      and _nfc(g.name).lower() == _nfc(f.name).lower()]
        if not candidatos:
            r.add("Continuidad", "—", f.name, "alerta", "no hay salida Python con ese nombre para comparar")
            continue
        g = max(candidatos, key=lambda x: x.stat().st_mtime)
        comp, detalle = _mismo_contenido(f, g)
        donde = f"{g.parent.parent.name}/{g.parent.name}"
        motivo = excepciones.get(_nfc(f.name).lower())
        if comp == "igual":
            r.add("Continuidad", g.parent.name, f.name, "ok", f"idéntico a la salida Python de {donde}")
        elif motivo is not None:
            r.add("Continuidad", g.parent.name, f.name, "alerta", f"diferencia conocida: {motivo} ({detalle})")
        else:
            r.add("Continuidad", g.parent.name, f.name, "alerta" if comp == "orden" else "error",
                  f"frente a la salida Python de {donde}: {detalle}")

    resumen = {e: sum(1 for i in r.items if i["estado"] == e) for e in ("ok", "alerta", "error")}
    out = {"periodo": periodo, "fecha": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
           "estado": "error" if resumen["error"] else ("alerta" if resumen["alerta"] else "ok"),
           "resumen": resumen, "modulos_activos": activos, "spa_activos": spa_activos,
           "verificaciones": r.items}
    if guardar and rep.is_dir():
        (rep / f"VALIDACION_{periodo}.json").write_text(json.dumps(out, ensure_ascii=False, indent=1),
                                                       encoding="utf-8")
        _guardar_excel(out, rep / f"VALIDACION_{periodo}.xlsx")
    return out


def _guardar_excel(out: dict, ruta: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    for fila in [["Período", out["periodo"]], ["Fecha", out["fecha"]], ["Estado", out["estado"].upper()],
                 ["OK", out["resumen"]["ok"]], ["Alertas", out["resumen"]["alerta"]],
                 ["Errores", out["resumen"]["error"]],
                 ["Módulos activos", ", ".join(out["modulos_activos"])],
                 ["Sin Precio Anterior", ", ".join(out["spa_activos"])]]:
        ws.append(fila)
    d = wb.create_sheet("Detalle")
    d.append(["Grupo", "Módulo", "Verificación", "Estado", "Detalle"])
    colores = {"ok": "C8E6C9", "alerta": "FFE0B2", "error": "FFCDD2"}
    for i in out["verificaciones"]:
        d.append([i["grupo"], i["modulo"], i["verificacion"], i["estado"].upper(), i["detalle"]])
        d.cell(d.max_row, 4).fill = PatternFill("solid", fgColor=colores[i["estado"]])
    for c in d[1]:
        c.font = Font(bold=True)
    for col, ancho in zip("ABCDE", (14, 12, 48, 9, 110)):
        d.column_dimensions[col].width = ancho
    wb.save(ruta)


def texto_resumen(out: dict) -> list[str]:
    icono = {"ok": "✔", "alerta": "⚠", "error": "✖"}
    lineas = [f"VALIDACIÓN {out['periodo']} — {icono[out['estado']]} {out['estado'].upper()} | "
              f"{out['resumen']['ok']} ok, {out['resumen']['alerta']} alertas, {out['resumen']['error']} errores"]
    for i in out["verificaciones"]:
        if i["estado"] != "ok":
            lineas.append(f"  {icono[i['estado']]} [{i['grupo']}] {i['modulo']} · {i['verificacion']}: {i['detalle']}")
    return lineas
